"""
Contains functions to easily edit the style of an element.
- Cell borders
- Table borders
- Font
- Paragraph
- TextFrame
"""

from .xml import SubElement
from pptx.dml.color import RGBColor as RGB
from pptx.enum.dml import MSO_LINE
from pptx.util import Pt


def set_cell_borders(cell,
                     border_color=RGB(0, 0, 0),
                     border_width=Pt(1),
                     border_style=MSO_LINE.DASH,
                     left=False,
                     top=False,
                     right=False,
                     bot=False):
    """
    The former border is kept (when adding a border, the former is removed if it exists)
    - border_color is the color of the border
    - border_width is the width of the border
    - border_style is the style of the border
    - left, top, right, bot indicate which borders are to be set
    """

    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()

    borders = [('L', left), ('R', right), ('T', top), ('B', bot)]
    for idx, (side, do) in enumerate(borders):
        if not do:
            continue
        tag = f"a:ln{side}"

        # Remove previous border if it exists
        tcPr.remove_all(tag)

        # Borders have to be declared in a specific order (L, R, T, B)
        # or they're not displayed
        before = [f"a:ln{s}" for (s, _) in borders[idx + 1:]]
        # borders must appear in the xml before the cell style
        # or they're not displayed
        before += ['a:solidFill', 'a:noFill']

        ln = SubElement(tcPr,
                        tag,
                        before=before,
                        w=str(border_width),
                        cap='flat',
                        cmpd='sng',
                        algn='ctr')

        solidFill = SubElement(ln, 'a:solidFill')
        # rgb color:
        colorval = "".join(map(lambda v: f"{v:02X}", border_color))
        SubElement(solidFill, 'a:srgbClr', val=colorval)
        # predefined color:
        # SubElement(solidFill, 'a:schemeClr', val="tx1")

        SubElement(ln, 'a:prstDash', val=MSO_LINE.to_xml(border_style))
        SubElement(ln, 'a:round')
        SubElement(ln, 'a:headEnd', type='none', w='med', len='med')
        SubElement(ln, 'a:tailEnd', type='none', w='med', len='med')


def set_rect_borders(table,
                     left,
                     top,
                     right,
                     bot,
                     outer=False,
                     inner=False,
                     **kwargs):
    """
    - left, top, right, bot are the coordinates of the border cells
    - outer indicates whether to set outer borders
    - inner indicates whether to set inner borders
    - kwargs are propagated to the Cell borders
    """

    assert left >= 0
    assert top >= 0
    assert right < len(table.columns)
    assert bot < len(table.rows)
    assert left <= right
    assert top <= bot

    for x in range(left, right + 1):
        for y in range(top, bot + 1):
            topb = outer if y == top else inner
            botb = outer if y == bot else inner
            leftb = outer if x == left else inner
            rightb = outer if x == right else inner

            set_cell_borders(table.cell(y, x),
                             **kwargs,
                             left=leftb,
                             top=topb,
                             right=rightb,
                             bot=botb)

    if not outer:
        return

    # PowerPoint has a bug (?), for cells strictly inside an array
    # only right and bottom borders are looked at, so we need to add
    # borders on cells outside of the selection

    for y in range(top, bot + 1):
        # right border of the left column
        if left != 0:
            set_cell_borders(table.cell(y, left - 1), **kwargs, right=True)
        # left border of the right column
        if right != len(table.columns) - 1:
            set_cell_borders(table.cell(y, right + 1), **kwargs, left=True)

    for x in range(left, right + 1):
        # lower border of upper row
        if top != 0:
            set_cell_borders(table.cell(top - 1, x), **kwargs, bot=True)
        # upper border of lower row
        if bot != len(table.rows) - 1:
            set_cell_borders(table.cell(bot + 1, x), **kwargs, top=True)


def font_style(ft,
               size=None,
               font=None,
               bold=None,
               italic=None,
               font_rgb=None,
               font_brightness=None,
               underline=None,
               language=None):
    """
    - size is the font size, usually Pt or Cm:
    https://python-pptx.readthedocs.io/en/latest/api/util.html
    - font is the name of the font (string, case sensitive)
    - bold indicates whether the text is bold
    - italic indicates whether the text is italic
    - font_rgb is the color of the text, of type RGBColor, built either from 3 integers
    or from a string corresponding to 3 hexadecimal integers:
    https://python-pptx.readthedocs.io/en/latest/api/dml.html#pptx.dml.color.RGBColor
    - font_brightness is the brightness of the font, a float between -1 and 1
    - underline indicates whether the text is underlined
    - language indicates the language of the text:
    https://python-pptx.readthedocs.io/en/latest/api/enum/MsoLanguageId.html
    """

    if size is not None:
        ft.size = size
    if font is not None:
        ft.name = font
    if bold is not None:
        ft.bold = bold
    if italic is not None:
        ft.italic = italic
    if font_rgb is not None:
        ft.color.rgb = font_rgb
    if font_brightness is not None:
        ft.color.brightness = font_brightness
    if underline is not None:
        ft.underline = underline
    if language is not None:
        ft.language_id = language


def text_style(tf, halign=None, **kwargs):
    """
    - halign indicates how the text is horizontally aligned
    https://python-pptx.readthedocs.io/en/latest/api/enum/PpParagraphAlignment.html
    - kwargs are propagated to the Paragraph
    """

    for paragraph in tf.paragraphs:
        font_style(paragraph.font, **kwargs)
        if halign is not None:
            paragraph.alignment = halign


def cell_style(cell,
               valign=None,
               back_rgb=None,
               back_brightness=None,
               margin_bottom=None,
               margin_top=None,
               margin_left=None,
               margin_right=None,
               **kwargs):
    """
    - valign indicated how the text is vertically aligned
    https://python-pptx.readthedocs.io/en/latest/api/enum/MsoVerticalAnchor.html
    - back_rgb is the background color of the cell
    https://python-pptx.readthedocs.io/en/latest/api/dml.html#pptx.dml.color.RGBColor
    - back_brightness is the background brightness of the cell, a float between -1 and 1
    - margin_top, margin_left, margin_bottom, margin_right are the margins of the cell
    https://python-pptx.readthedocs.io/en/latest/api/util.html
    - kwargs are propagated to the TextFrame
    """

    if margin_bottom is not None:
        cell.margin_bottom = margin_bottom
    if margin_top is not None:
        cell.margin_top = margin_top
    if margin_right is not None:
        cell.margin_right = margin_right
    if margin_left is not None:
        cell.margin_left = margin_left
    if valign is not None:
        cell.vertical_anchor = valign
    if back_brightness is not None:
        cell.fill.solid()
        cell.fill.fore_color.brightness = back_brightness
    if back_rgb is not None:
        if not back_rgb:
            cell.fill.background()
        else:
            cell.fill.solid()
            cell.fill.fore_color.rgb = back_rgb

    text_style(cell.text_frame, **kwargs)


def paragraph_append(paragraph, txt, **kwargs):
    """
    Adds a run to the paragraph, with the given text.

    - kwargs are propagated to the Run.
    """
    run = paragraph.add_run()
    run.text = txt
    font_style(run.font, **kwargs)
